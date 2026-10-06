# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import os

folder = r"G:\我的雲端硬碟\0_AI Agent\Obsidian\DavidCloud\大衛人生\B3. 健康養生\B3.1 精力體態\B3.1.3 健康可視化"
xlsx_path = os.path.join(folder, "米家八電極體脂數據追蹤表.xlsx")
csv_path = os.path.join(folder, "米家八電極體脂數據追蹤表.csv")

headers = [
    "日期", "體重 (kg)", "體脂率 (%)", "BMI", "心率 (bpm)", "綜合得分",
    "去脂體重 (kg)", "脂肪量 (kg)", "肌肉量 (kg)", "肌肉率 (%)", "骨骼肌量 (kg)", "SMI (kg/m²)",
    "蛋白質 (kg)", "蛋白質率 (%)", "水分 (kg)", "水分率 (%)", "骨鹽量 (kg)", "內臟脂肪等級",
    "基礎代謝率 (kcal)", "建議攝取熱量 (kcal)", "推估腰臀比",
    "軀幹脂肪 (kg)", "左臂脂肪 (kg)", "右臂脂肪 (kg)", "左腿脂肪 (kg)", "右腿脂肪 (kg)",
    "軀幹肌肉 (kg)", "左臂肌肉 (kg)", "右臂肌肉 (kg)", "左腿肌肉 (kg)", "右腿肌肉 (kg)",
    "距Q4體重差距 (kg)", "距Q4體脂差距 (%)", "備註 / 圖片來源"
]

data = [
    [
        "2026-10-03", 85.65, 29.0, 26.4, 93, 71,
        60.8, 24.8, 57.5, 67.2, 32.4, 7.9,
        12.0, 14.0, 44.9, 52.5, 3.3, 10.0,
        1683, 2440, 0.90,
        13.0, 1.7, 1.7, 3.6, 3.6,
        26.4, 3.1, 3.1, 9.6, 9.6,
        1.65, 3.0, "體脂/20261003.jpg (米家八電極 S800)"
    ],
    [
        "2026-10-04", 85.80, 29.8, 26.5, 96, 70,
        60.2, 25.6, 56.9, 66.3, 32.2, 7.7,
        11.9, 13.9, 44.4, 51.7, 3.3, 11.0,
        1670, 2422, 0.90,
        13.3, 1.8, 1.7, 3.7, 3.7,
        26.1, 3.0, 3.0, 9.5, 9.5,
        1.80, 3.8, "體脂/20261004.jpg (米家八電極 S800)"
    ]
]

# 產出 CSV
df = pd.DataFrame(data, columns=headers)
df.to_csv(csv_path, index=False, encoding="utf-8-sig")

# 產出美化 Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "每日體脂量測"

# 寫入表頭與資料
ws.append(headers)
for row in data:
    ws.append(row)

# 樣式設定
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
data_font = Font(name="微軟正黑體", size=10)
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

# 凍結首行與前兩欄
ws.freeze_panes = "C2"

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row_idx in range(2, len(data) + 2):
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = data_font
        cell.border = thin_border
        if col_idx == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx == len(headers):
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center")

# 自動調整欄寬
for col in ws.columns:
    max_len = 0
    for cell in col:
        val = str(cell.value or "")
        # 計算字元長度（中文字算 2 個字元）
        clen = sum(2 if ord(c) > 127 else 1 for c in val)
        if clen > max_len:
            max_len = clen
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

ws.row_dimensions[1].height = 28
ws.row_dimensions[2].height = 22

wb.save(xlsx_path)
print("SUCCESS: Generated XLSX and CSV successfully!")
